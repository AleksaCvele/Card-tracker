import re
from typing import List, Dict, Tuple, Optional, Any
from openpyxl import load_workbook
from collection import CollectionStorage


class CardImporter:

    @staticmethod
    def parse_universal_line(line: str) -> Optional[Tuple[int, str, Optional[str], Optional[str]]]:
        cleaned = line.strip()
        if not cleaned or cleaned.startswith("//") or cleaned.startswith("#"):
            return None

        pattern_full = r"^(\d+)\s+(.+?)\s+\(([A-Za-z0-9]+)\)\s*([A-Za-z0-9\-]+)?"
        match = re.match(pattern_full, cleaned)
        if match:
            qty = int(match.group(1))
            name = match.group(2).strip()
            set_code = match.group(3).strip().lower()
            coll_num = match.group(4).strip() if match.group(4) else None
            return qty, name, set_code, coll_num

        pattern_simple = r"^(\d+)\s+(.+)"
        match_simple = re.match(pattern_simple, cleaned)
        if match_simple:
            qty = int(match_simple.group(1))
            name = match_simple.group(2).strip()
            return qty, name, None, None

        return 1, cleaned, None, None

    @staticmethod
    def parse_image_uri(line_or_text: str) -> Optional[str]:
    
        pattern = r'"large"\s*:\s*"(https://[^"]+)"'
        match = re.search(pattern, line_or_text)
        return match.group(1) if match else None

    @staticmethod
    def match_cards_with_database(
        parsed_items: List[Tuple[int, str, Optional[str], Optional[str]]],
        all_cards_database: List[Any]
    ) -> List[Any]:
        db_exact_map: Dict[Tuple[str, str, str], Any] = {}
        for card in all_cards_database:
            c_dict = card.to_dict() if hasattr(card, "to_dict") else card
            key = (
                c_dict.get("name", "").lower(),
                c_dict.get("set", "").lower(),
                str(c_dict.get("collector_number", ""))
            )
            if key not in db_exact_map:
                db_exact_map[key] = card

        db_cheapest_map = CollectionStorage._find_cheapest_versions_map(all_cards_database)

        matched_cards = []
        for qty, name, set_code, coll_num in parsed_items:
            name_key = name.lower()
            found_card = None

            if set_code and coll_num:
                exact_key = (name_key, set_code.lower(), str(coll_num))
                found_card = db_exact_map.get(exact_key)

            if not found_card:
                found_card = db_cheapest_map.get(name_key)

            if found_card:
                if hasattr(found_card, "clone"):
                    card_clone = found_card.clone(quantity=qty)
                    card_clone.quantity = qty
                    matched_cards.append(card_clone)
                else:
                    card_copy = found_card.copy() if isinstance(found_card, dict) else dict(found_card)
                    card_copy["quantity"] = qty
                    matched_cards.append(card_copy)

        return matched_cards

    @classmethod
    def parse_card_list_text(cls, text_content: str, all_cards_database: List[Any]) -> List[Any]:
        parsed_items = []
        for line in text_content.splitlines():
            item = cls.parse_universal_line(line)
            if item:
                parsed_items.append(item)

        return cls.match_cards_with_database(parsed_items, all_cards_database)

    @classmethod
    def parse_excel_file(cls, filepath: str, all_cards_database: List[Any]) -> List[Any]:
        try:
            wb = load_workbook(filename=filepath, data_only=True)
            sheet = wb.active
        except Exception as e:
            raise ValueError(f"Greška pri otvaranju Excel fajla '{filepath}': {e}")

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        
        col_name = cls._find_column_index(header, ["name", "card name", "naziv", "kartica"])
        col_qty = cls._find_column_index(header, ["quantity", "qty", "kolicina", "count", "count/qty"])
        col_set = cls._find_column_index(header, ["set", "set code", "set_code", "edition"])
        col_num = cls._find_column_index(header, ["collector number", "collector_number", "number", "card #", "#"])

        if col_name is None:
            raise ValueError("Excel fajl mora sadržati kolonu sa nazivom kartice ('Name' ili 'Naziv').")

        parsed_items = []
        for row in rows[1:]:
            if not row or row[col_name] is None:
                continue

            name = str(row[col_name]).strip()
            if not name:
                continue

            qty = 1
            if col_qty is not None and row[col_qty] is not None:
                try:
                    qty = int(row[col_qty])
                except (ValueError, TypeError):
                    qty = 1

            set_code = None
            if col_set is not None and row[col_set] is not None:
                set_code = str(row[col_set]).strip().lower()

            coll_num = None
            if col_num is not None and row[col_num] is not None:
                coll_num = str(row[col_num]).strip()

            parsed_items.append((qty, name, set_code, coll_num))

        return cls.match_cards_with_database(parsed_items, all_cards_database)

    @staticmethod
    def _find_column_index(header: List[str], possible_names: List[str]) -> Optional[int]:
        for name in possible_names:
            if name in header:
                return header.index(name)
        return None

def parse_universal_line(line: str):
    return CardImporter.parse_universal_line(line)

def parse_image_uri(line_or_text: str):
    return CardImporter.parse_image_uri(line_or_text)

def match_cards_with_database(parsed_items, all_cards_database):
    return CardImporter.match_cards_with_database(parsed_items, all_cards_database)

def parse_card_list_text(text_content: str, all_cards_database: List[Any]):
    return CardImporter.parse_card_list_text(text_content, all_cards_database)

def parse_excel_file(filepath: str, all_cards_database: List[Any]):
    return CardImporter.parse_excel_file(filepath, all_cards_database)