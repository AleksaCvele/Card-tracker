import json
import os
from typing import List, Tuple, Optional, Dict, Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class CollectionExporter:

    @staticmethod
    def export_to_excel(cards: List[Any], filename: str = "moja_kolekcija.xlsx") -> Tuple[bool, str]:
        
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Kolekcija"

            # Zaglavlja kolona
            headers = [
                "Naziv kartice",
                "Set iz kog je kartica",
                "Kolekcijski broj kartice",
                "Cena kartice"
            ]
            ws.append(headers)

            # Stil za zaglavlje
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            # Upis podataka
            for card in cards:
                if hasattr(card, "name"):
                    name = card.name
                    set_name = card.set_name
                    collector_num = str(card.collector_number)
                    price = card.price_normal if card.price_normal != "N/A" else card.price_foil
                else:
                    name = card.get("name", "")
                    set_name = card.get("set", card.get("set_name", ""))
                    collector_num = str(card.get("collector_number", ""))
                    price = card.get("price_normal", "N/A")
                    if price == "N/A":
                        price = card.get("price_foil", "N/A")

                ws.append([name, set_name, collector_num, price])

            # Poravnanje i ivice ćelija sa podacima
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
                row[0].alignment = Alignment(horizontal="left", vertical="center")      # Naziv
                row[1].alignment = Alignment(horizontal="left", vertical="center")      # Set
                row[2].alignment = Alignment(horizontal="center", vertical="center")    # Kolekcijski broj
                row[3].alignment = Alignment(horizontal="right", vertical="center")     # Cena

                for cell in row:
                    cell.border = thin_border

            # Pristajanje širine kolona (100% vidljivost teksta)
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)

                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(filename)
            return True, filename

        except Exception as e:
            print(f"Greška pri izvozu u Excel: {e}")
            return False, filename


class CollectionStorage:

    @classmethod
    def save_to_file(cls, filename: str, collection_list: List[Any]):
        temp_filename = filename + ".tmp"
        try:
            saved_data = []
            for card in collection_list:
                card_dict = card.to_dict() if hasattr(card, "to_dict") else card
                saved_data.append({
                    "name": card_dict.get("name", ""),
                    "set": card_dict.get("set", card_dict.get("set_name", "")),
                    "collector_number": str(card_dict.get("collector_number", "")),
                    "quantity": card_dict.get("quantity", 1),
                    "is_foil": card_dict.get("is_foil", False),
                    "cmc": card_dict.get("cmc", 0),
                    "colors": card_dict.get("colors", ""),
                    "type": card_dict.get("type", card_dict.get("type_line", "")),
                    "rarity": card_dict.get("rarity", ""),
                    "price_normal": card_dict.get("price_normal", "N/A"),
                    "price_foil": card_dict.get("price_foil", "N/A"),
                    "price_numeric": card_dict.get("price_numeric", 0.0),
                    "price_rsd": card_dict.get("price_rsd", 0.0)
                })

            with open(temp_filename, "w", encoding="utf-8") as f:
                json.dump(saved_data, f, ensure_ascii=False, indent=4)

            os.replace(temp_filename, filename)
            
        except Exception as e:
            print(f"Greška pri čuvanju kolekcije: {e}")
            if os.path.exists(temp_filename):
                try:
                    os.remove(temp_filename)
                except OSError:
                    pass
            raise e

    @classmethod
    def load_from_file(cls, filename: str, all_cards_database: List[Any]) -> List[Any]:
      
        if not os.path.exists(filename):
            # Provera ako je prošireno bez ekstenzije
            if os.path.exists(filename + ".xlsx"):
                filename += ".xlsx"
            elif os.path.exists(filename + ".json"):
                filename += ".json"
            else:
                return []

        if filename.endswith(".xlsx"):
            return cls._load_from_excel(filename, all_cards_database)
        
        return cls._load_from_json(filename, all_cards_database)

    @classmethod
    def _load_from_json(cls, filename: str, all_cards_database: List[Any]) -> List[Any]:
        db_map, db_cheapest_map = cls._build_database_maps(all_cards_database)
        loaded_collection = []

        try:
            with open(filename, "r", encoding="utf-8") as f:
                saved_data = json.load(f)

            for item in saved_data:
                exact_key = (
                    item.get("name", "").lower(),
                    item.get("set", "").lower(),
                    str(item.get("collector_number", ""))
                )
                name_key = item.get("name", "").lower()
                qty = item.get("quantity", 1)
                is_foil = item.get("is_foil", False)

                found_card = None
                if exact_key in db_map:
                    found_card = db_map[exact_key]
                elif name_key in db_cheapest_map:
                    found_card = db_cheapest_map[name_key]

                if found_card:
                    if hasattr(found_card, "clone"):
                        card_clone = found_card.clone(is_foil=is_foil, quantity=qty)
                        card_clone.quantity = qty
                        loaded_collection.append(card_clone)
                    else:
                        card_copy = found_card.copy() if isinstance(found_card, dict) else dict(found_card)
                        card_copy["quantity"] = qty
                        card_copy["is_foil"] = is_foil
                        loaded_collection.append(card_copy)

        except Exception as e:
            print(f"Greška pri učitavanju JSON kolekcije: {e}")

        return loaded_collection

    
    @classmethod
    def _load_from_excel(cls, filename: str, all_cards_database: List[Any]) -> List[Any]:
        try:
            wb = openpyxl.load_workbook(filename=filename, data_only=True)
            sheet = wb.active
        except Exception as e:
            print(f"Greška pri čitanju Excel fajla '{filename}': {e}")
            return []

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        # Detekcija indeksa kolona iz zaglavlja (prvi red)
        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        
        col_name = cls._find_column_index(header, ["naziv kartice", "name", "card name", "naziv"])
        col_set = cls._find_column_index(header, ["set iz kog je kartica", "set", "set code", "edition"])
        col_num = cls._find_column_index(header, ["kolekcijski broj kartice", "collector number", "collector_number", "number", "card #", "#"])
        col_qty = cls._find_column_index(header, ["kolicina", "quantity", "qty", "count"])

        if col_name is None:
            print("Excel fajl ne sadrži prepoznatljivu kolonu sa nazivom kartice.")
            return []

        db_map, db_cheapest_map = cls._build_database_maps(all_cards_database)
        loaded_collection = []

        for row in rows[1:]:
            if not row or row[col_name] is None:
                continue

            name = str(row[col_name]).strip()
            if not name:
                continue

            set_code = str(row[col_set]).strip().lower() if col_set is not None and row[col_set] is not None else ""
            coll_num = str(row[col_num]).strip() if col_num is not None and row[col_num] is not None else ""
            
            # Ako u Excelu postoji kolona za količinu, uzimamo taj broj, inače 1
            qty = 1
            if col_qty is not None and row[col_qty] is not None:
                try:
                    qty = int(row[col_qty])
                except (ValueError, TypeError):
                    qty = 1

            exact_key = (name.lower(), set_code, coll_num)
            name_key = name.lower()

            found_card = None
            if set_code and coll_num and exact_key in db_map:
                found_card = db_map[exact_key]
            elif name_key in db_cheapest_map:
                found_card = db_cheapest_map[name_key]

            if found_card:
                for _ in range(qty):
                    loaded_collection.append(found_card)

        return loaded_collection

    @staticmethod
    def _build_database_maps(all_cards_database: List[Any]) -> Tuple[Dict, Dict]:
        db_map = {}
        for card in all_cards_database:
            c_dict = card.to_dict() if hasattr(card, "to_dict") else card
            key = (
                c_dict.get("name", "").lower(),
                c_dict.get("set", c_dict.get("set_name", "")).lower(),
                str(c_dict.get("collector_number", ""))
            )
            if key not in db_map:
                db_map[key] = card

        db_cheapest_map = CollectionStorage._find_cheapest_versions_map(all_cards_database)
        return db_map, db_cheapest_map

    @staticmethod
    def _find_column_index(header: List[str], possible_names: List[str]) -> Optional[int]:
        for name in possible_names:
            if name in header:
                return header.index(name)
        return None

    @staticmethod
    def _find_cheapest_versions_map(all_cards: List[Any]) -> Dict[str, Any]:
        cheapest_map = {}
        for card in all_cards:
            c_dict = card.to_dict() if hasattr(card, "to_dict") else card
            name_key = c_dict.get("name", "").lower()
            price = c_dict.get("price_numeric", 0.0)

            if name_key not in cheapest_map:
                cheapest_map[name_key] = card
            else:
                existing_card = cheapest_map[name_key]
                ex_dict = existing_card.to_dict() if hasattr(existing_card, "to_dict") else existing_card
                current_cheapest_price = ex_dict.get("price_numeric", 0.0)

                if price > 0 and (current_cheapest_price == 0 or price < current_cheapest_price):
                    cheapest_map[name_key] = card

        return cheapest_map

def save_collection_to_file(collection_list: List[Any], filename: str) -> Tuple[bool, str]:
    return CollectionStorage.save_to_file(collection_list, filename)


def load_collection_from_file(filename: str, all_cards_database: List[Any]) -> List[Any]:
    return CollectionStorage.load_from_file(filename, all_cards_database)


def export_collection_to_excel(collection_list: List[Any], filename: str = "kolekcija.xlsx") -> Tuple[bool, str]:
    return CollectionExporter.export_to_excel(collection_list, filename)